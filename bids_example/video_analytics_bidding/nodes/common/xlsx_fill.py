"""Filling the buyer's spreadsheet templates.

The RFP asks for commercials and sizing in the buyer's own formats, so the shipped
templates are loaded and edited in place -- never rebuilt. `openpyxl` preserves cell
styling, column widths and merged ranges across a load/save round-trip, but only if the
workbook is opened rather than recreated.

Two mechanics this module exists to get right:

**Merged cells.** Both templates merge heavily -- the sizing totals are single cells
merged down twenty rows (`E4:E23`), and every commercial price cell is merged with its
neighbour (`H10:I10`). Only the top-left anchor of a merged range is writable; assigning
to any other cell in the range raises `AttributeError` because openpyxl represents them
as read-only `MergedCell`s. `set_cell` resolves to the anchor so callers can address a
range by any of its coordinates.

**Formulas.** The templates carry live formulas (`=(F10*D10)`, `=SUM(C5:C23)`,
`=ROUND(D4/60,0)`). openpyxl does not evaluate them -- reading such a cell returns the
formula string, not a number. So a caller can never read a total back out of the
workbook, and any number the agent must *state* (the four sizing totals, the commercial
aggregates) is computed in Python and written as a literal. Formulas that only serve the
human reading the delivered file are left alone, so the document still recalculates when
opened.
"""
import logging
import os
import tempfile

import openpyxl
from openpyxl.utils import get_column_letter

log = logging.getLogger(__name__)

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_DIR = os.path.abspath(os.path.join(HERE, ".."))

COMMERCIALS_TEMPLATE = os.path.join(TEMPLATE_DIR, "Commercials_Template_formated.xlsx")
SIZING_TEMPLATE = os.path.join(TEMPLATE_DIR, "Sizing_Template_simple_formated.xlsx")


def load(template_path):
    """Open a template for editing. Formulas are kept as formulas."""
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"spreadsheet template not found: {template_path}")
    return openpyxl.load_workbook(template_path)


def merged_anchor(ws, coordinate):
    """The writable top-left cell of whatever merged range `coordinate` falls in.

    Returns `coordinate` unchanged when the cell is not merged.
    """
    for rng in ws.merged_cells.ranges:
        if coordinate in rng:
            return rng.start_cell.coordinate
    return coordinate


def set_cell(ws, coordinate, value):
    """Write one cell, resolving merged ranges to their anchor."""
    anchor = merged_anchor(ws, coordinate)
    if anchor != coordinate:
        log.debug("%s is merged; writing to anchor %s", coordinate, anchor)
    ws[anchor] = value
    return anchor


def set_cells(ws, mapping):
    """Write a `{coordinate: value}` mapping."""
    for coordinate, value in mapping.items():
        set_cell(ws, coordinate, value)


def fill_rows(ws, start_row, column_map, records, clear_to_row=None):
    """Write `records` down the sheet from `start_row`.

    `column_map` maps a record key to a column letter, e.g.
    `{"usecase": "C", "licenses": "D", "unit_price": "F"}`.

    Rows the records do not reach, up to `clear_to_row`, have their mapped columns
    blanked. Without that, a template pre-populated with seventeen sample use cases
    would silently leave the leftovers in a bid that only quotes four -- the buyer would
    read priced lines the company never offered.

    Returns the list of row numbers written.
    """
    written = []
    row = start_row
    for record in records:
        for key, column in column_map.items():
            if key in record:
                set_cell(ws, f"{column}{row}", record[key])
        written.append(row)
        row += 1

    if clear_to_row:
        for blank_row in range(row, clear_to_row + 1):
            for column in column_map.values():
                set_cell(ws, f"{column}{blank_row}", None)
        if row <= clear_to_row:
            log.info("cleared unused template rows %d-%d", row, clear_to_row)

    return written


def save_temp(wb, prefix):
    """Save to a temp .xlsx and return the path. The caller uploads and removes it."""
    fd, path = tempfile.mkstemp(suffix=".xlsx", prefix=prefix)
    os.close(fd)
    wb.save(path)
    log.info("wrote workbook %s", path)
    return path


def describe(ws, max_row=30, max_col=11):
    """Non-empty cells as `{coordinate: value}`. Used by tests and for debugging."""
    out = {}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_row),
                            max_col=min(ws.max_column, max_col)):
        for cell in row:
            if cell.value is not None:
                out[cell.coordinate] = cell.value
    return out
